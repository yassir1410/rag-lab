import io
import json
import shutil
import tempfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import docx
import fitz  # PyMuPDF

from app.core.config import settings
from app.db.base import Base
from app.db.session import get_db
from app.main import app

# In-memory SQLite engine for fast, isolated tests
TEST_SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

test_engine = create_engine(
    TEST_SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
    json_serializer=lambda obj: json.dumps(obj, default=str),
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


@pytest.fixture(scope="function")
def db_session():
    """Create a fresh database session for each test function."""
    Base.metadata.create_all(bind=test_engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=test_engine)


@pytest.fixture(scope="function")
def client(db_session, monkeypatch):
    """FastAPI test client with database and storage overridden."""
    # Create isolated temp directory for uploads
    temp_dir = tempfile.mkdtemp()
    temp_path = Path(temp_dir)
    monkeypatch.setattr(settings, "STORAGE_DIR", temp_dir)

    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()
    shutil.rmtree(temp_dir, ignore_errors=True)


# ==========================================
# Synthetic Test Document Fixtures
# ==========================================

@pytest.fixture
def synthetic_txt_bytes() -> bytes:
    """UTF-8 text with multiple paragraphs."""
    content = (
        "Introduction to Retrieval-Augmented Generation.\n\n"
        "RAG enhances large language models by combining parametric knowledge with real-time dynamic retrieval.\n\n"
        "This ensures answers are grounded in private and up-to-date documentation."
    )
    return content.encode("utf-8")


@pytest.fixture
def synthetic_txt_bom_bytes() -> bytes:
    """UTF-8 text with BOM signature."""
    content = "Document with Byte Order Mark.\n\nValidating encoding resilience."
    return content.encode("utf-8-sig")


@pytest.fixture
def synthetic_md_bytes() -> bytes:
    """Markdown document with headings, paragraphs, code fences, and tables."""
    content = """# System Architecture

RAG Engineering Lab is built as a modular monorepo.

## Code Example

Here is a Python function:
```python
def compute_similarity(v1, v2):
    return sum(a * b for a, b in zip(v1, v2))
```

## Performance Metrics

| Metric | Target | Actual |
| --- | --- | --- |
| Latency | 50ms | 22ms |
| Precision | 0.90 | 0.94 |
"""
    return content.encode("utf-8")


@pytest.fixture
def synthetic_pdf_bytes() -> bytes:
    """Multi-page PDF document generated dynamically with PyMuPDF."""
    doc = fitz.open()

    # Page 1: Heading and text
    page1 = doc.new_page()
    page1.insert_text(
        fitz.Point(50, 72),
        "Chapter 1: Principles of Retrieval\n\nRetrieval systems locate relevant context from dense vector stores.",
        fontsize=12,
    )

    # Page 2: Second page
    page2 = doc.new_page()
    page2.insert_text(
        fitz.Point(50, 72),
        "Chapter 2: Context Synthesis\n\nAugmenting prompts prevents factual hallucinations.",
        fontsize=12,
    )

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def synthetic_empty_pdf_bytes() -> bytes:
    """PDF with a blank page (no text)."""
    doc = fitz.open()
    doc.new_page()
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def synthetic_docx_bytes() -> bytes:
    """DOCX document with headings, paragraphs, and tables generated dynamically."""
    doc = docx.Document()
    doc.add_heading("Deep Learning Foundations", level=1)
    doc.add_paragraph("Neural networks learn representations through backpropagation.")

    doc.add_heading("Model Parameters", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Parameter"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Batch Size"
    table.rows[1].cells[1].text = "32"

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
